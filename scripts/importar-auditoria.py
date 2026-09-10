"""Converte a auditoria Excel em dados leves para o painel do site."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl


def clean(value):
    if value in (None, "", "-"):
        return None
    return value


def text(value, fallback="Não informado"):
    value = clean(value)
    if value is None:
        return fallback
    result = str(value).strip()
    return "Não informado" if result.casefold() == "não informado" else result


def status_projeto(value):
    value = text(value)
    upper = value.upper()
    if "RETIRADO" in upper:
        return "Retirado"
    if "ARQUIVADO" in upper:
        return "Arquivado"
    return value


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("entrada", type=Path)
    parser.add_argument("saida", type=Path)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.entrada, data_only=True, read_only=True)
    recursos = []
    legislacao = []

    oficiais = workbook["4. Emendas Internas"]
    for index, row in enumerate(oficiais.iter_rows(min_row=7, max_row=29, values_only=True), start=1):
        ano, ficha, valor, destino, categoria, hospital, responsavel, execucao, ano_execucao, onde, observacao, fonte, link = row
        recursos.append({
            "id": f"oficial-{index}",
            "ano": text(ano),
            "ficha": text(ficha),
            "valor": float(valor),
            "destino": text(destino),
            "categoria": text(categoria),
            "hospital": text(hospital),
            "responsavel": text(responsavel),
            "comprovacao": "Oficial",
            "execucao": text(execucao),
            "anoExecucao": text(ano_execucao),
            "observacao": text(observacao, ""),
            "fonte": text(fonte),
            "link": text(link, ""),
        })

    informais = workbook["6. Emendas Informais"]
    for index, row in enumerate(informais.iter_rows(min_row=6, max_row=8, values_only=True), start=1):
        ano, valor, destino, categoria, hospital, observacao, verificar, fonte = row
        recursos.append({
            "id": f"informal-valor-{index}",
            "ano": text(ano),
            "ficha": "Não informada",
            "valor": float(valor),
            "destino": text(destino),
            "categoria": text(categoria),
            "hospital": text(hospital),
            "responsavel": "Atribuído a Kiel no material de campanha",
            "comprovacao": "Informal",
            "execucao": "Não confirmada",
            "anoExecucao": "Não informado",
            "observacao": text(observacao, ""),
            "verificar": text(verificar, ""),
            "fonte": text(fonte),
            "link": "",
        })

    for index, row in enumerate(informais.iter_rows(min_row=13, max_row=25, values_only=True), start=1):
        item, categoria, observacao, _, _, verificar, _, fonte = row
        recursos.append({
            "id": f"informal-sem-valor-{index}",
            "ano": "Não informado",
            "ficha": "Não informada",
            "valor": None,
            "destino": text(item),
            "categoria": text(categoria),
            "hospital": "Não informado",
            "responsavel": "Atribuído a Kiel no material de campanha",
            "comprovacao": "Informal",
            "execucao": "Não confirmada",
            "anoExecucao": "Não informado",
            "observacao": text(observacao, ""),
            "verificar": text(verificar, ""),
            "fonte": text(fonte),
            "link": "",
        })

    leis = workbook["2. Leis Aprovadas"]
    for index, row in enumerate(leis.iter_rows(min_row=7, max_row=21, values_only=True), start=1):
        numero, aprovacao, apresentacao, pl_origem, ementa, responsaveis, mandato, status, fonte, verificar = row
        legislacao.append({
            "id": f"lei-{index}",
            "tipo": "Lei ordinária",
            "numero": text(numero),
            "ano": text(aprovacao).split("/")[-1],
            "data": text(aprovacao),
            "plOrigem": text(pl_origem),
            "descricao": text(ementa),
            "responsaveis": text(responsaveis),
            "mandato": text(mandato),
            "status": text(status).capitalize(),
            "fonte": text(fonte),
            "verificar": text(verificar, ""),
        })

    projetos = workbook["3. Projetos Não Aprovados"]
    for index, row in enumerate(projetos.iter_rows(min_row=6, max_row=14, values_only=True), start=1):
        numero, entrada, ementa, responsaveis, desfecho, data_desfecho, fonte, link = row
        ano = text(numero).split("/")[-1]
        legislacao.append({
            "id": f"projeto-{index}",
            "tipo": "Projeto de lei",
            "numero": text(numero),
            "ano": ano,
            "data": text(entrada),
            "plOrigem": text(numero),
            "descricao": text(ementa),
            "responsaveis": text(responsaveis),
            "mandato": "1º mandato (2013–2016)" if int(ano) <= 2016 else "3º mandato (2021–2024)",
            "status": status_projeto(desfecho),
            "desfecho": text(desfecho),
            "dataDesfecho": text(data_desfecho),
            "fonte": text(fonte),
            "link": text(link, ""),
        })

    data = {
        "meta": {
            "titulo": "Atuação e recursos — Kiel Damasceno",
            "atualizacao": "setembro de 2026",
            "coberturaLeis": "2013 a 2024",
            "coberturaRecursos": "2022 ao 1º trimestre de 2025",
            "arquivoOrigem": args.entrada.name,
        },
        "recursos": recursos,
        "legislacao": legislacao,
    }
    args.saida.parent.mkdir(parents=True, exist_ok=True)
    content = "window.KIEL_ATUACAO_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
    args.saida.write_text(content, encoding="utf-8")
    print(f"{len(recursos)} recursos e {len(legislacao)} registros legislativos exportados para {args.saida}")


if __name__ == "__main__":
    main()
